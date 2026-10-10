import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_page_number_to_footer(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.name = "Times New Roman"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(100, 116, 139)
    
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)
    
    run2 = paragraph.add_run(" of ")
    run2.font.name = "Times New Roman"
    run2.font.size = Pt(9.5)
    run2.font.color.rgb = RGBColor(100, 116, 139)
    
    run3 = paragraph.add_run()
    run3.font.name = "Times New Roman"
    run3.font.size = Pt(9.5)
    run3.font.color.rgb = RGBColor(100, 116, 139)
    fldChar4 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText2 = parse_xml(r'<w:instrText %s xml:space="preserve"> NUMPAGES </w:instrText>' % nsdecls('w'))
    fldChar5 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar6 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run3._r.append(fldChar4)
    run3._r.append(instrText2)
    run3._r.append(fldChar5)
    run3._r.append(fldChar6)

def style_paragraph(p, space_before=0, space_after=4, line_spacing=1.15):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    style_paragraph(p, space_before=6, space_after=6, line_spacing=1.15)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    style_paragraph(p, space_before=4, space_after=3, line_spacing=1.15)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 58, 138) # Navy Blue
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    style_paragraph(p, space_before=3, space_after=2, line_spacing=1.15)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = RGBColor(51, 65, 85) # Slate 700
    return p

def add_body_p(doc, text, bold_prefix=None, space_after=4, italic_prefix=None):
    p = doc.add_paragraph()
    style_paragraph(p, space_before=0, space_after=space_after, line_spacing=1.15)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
    if italic_prefix:
        r_pre = p.add_run(italic_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(10.5)
        r_pre.font.italic = True
        r_pre.font.color.rgb = RGBColor(30, 58, 138)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_bullet_p(doc, bold_prefix, text, space_after=3):
    p = doc.add_paragraph(style='List Bullet')
    style_paragraph(p, space_before=0, space_after=space_after, line_spacing=1.15)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_callout(doc, text, title="KEY ARCHITECTURAL HIGHLIGHT:"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.25)
    cell = table.cell(0, 0)
    set_cell_background(cell, "EFF6FF")
    set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
    p = cell.paragraphs[0]
    style_paragraph(p, space_before=0, space_after=2, line_spacing=1.12)
    r_t = p.add_run(title + " ")
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(10)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(30, 58, 138)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(30, 41, 59)
    p_dummy = doc.add_paragraph()
    style_paragraph(p_dummy, space_before=0, space_after=2, line_spacing=1.0)

def add_code_block(doc, code_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.25)
    cell = table.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=70, bottom=70, left=120, right=120)
    p = cell.paragraphs[0]
    style_paragraph(p, space_before=0, space_after=0, line_spacing=1.05)
    r = p.add_run(code_text)
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(15, 23, 42)
    p_dummy = doc.add_paragraph()
    style_paragraph(p_dummy, space_before=0, space_after=2, line_spacing=1.0)

def add_figure_image(doc, img_path, caption_text, width_in=5.4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, space_before=2, space_after=2, line_spacing=1.0)
    p.add_run().add_picture(img_path, width=Inches(width_in))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p_cap, space_before=1, space_after=4, line_spacing=1.05)
    r_cap = p_cap.add_run(caption_text)
    r_cap.font.name = "Times New Roman"
    r_cap.font.size = Pt(9.5)
    r_cap.font.bold = True
    r_cap.font.color.rgb = RGBColor(30, 58, 138)

def format_custom_table(table, header_bg="1E3A8A", alt_bg="F1F5F9"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, col in enumerate(table.columns):
        pass
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            if r_idx == 0:
                set_cell_background(cell, header_bg)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    style_paragraph(p, space_before=1, space_after=1, line_spacing=1.0)
                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(9.5)
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
            else:
                bg = alt_bg if r_idx % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg)
                for p in cell.paragraphs:
                    style_paragraph(p, space_before=1, space_after=1, line_spacing=1.05)
                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(9)
                        r.font.color.rgb = RGBColor(30, 41, 59)

print("Base helper routines loaded successfully.")
